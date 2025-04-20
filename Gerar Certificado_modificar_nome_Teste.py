#Instalar biblioteca pip install python-docx
#Documentação: https://python-docx.readthedocs.io/en/latest/

# Realiza a modificação manualmente, inserindo informações do aluno

from docx import Document # Permite abrir e mexem em arquivos .docx
from docx.shared import Pt # com tamanho da fonte em "pontos", tipo 12 pt, igual no Word

from openpyxl import load_workbook #Abrir planilhas do excel .xlsx
import os # Sistema Operacional, caminho de arquivos, ainda não utilizada no código

alunos = "Alunos.xlsx"
planilha_alunos = load_workbook(alunos) # Armazena na variavel a planilha anexada na pasta, alunos.xlsx

sheet_selecionada = planilha_alunos["Nomes"] # Aba da planilha

for linha in range(2, len(sheet_selecionada["A"])+ 1): # Range(2, ...), começar a partir da linha 2 no excel, e len(..., conta
    # quantas células tem na coluna A, onde estão os nomes

    word = Document("Certificado1.docx")#Modelo exemplo, vai abrir

    estilo = word.styles["Normal"] #é o que define tipo de letra, tamanho, etc.

#um "marcador" A, vai pagar o conteúdo da coluna A, linha por linha
    nome_aluno = sheet_selecionada[f'A{linha}'].value

    for paragrafo in word.paragraphs: #Aqui ele está olhando parágrafo por parágrafo no documento do Word, procurando o texto @nome.
        if "@nome" in paragrafo.text:
            paragrafo.text = nome_aluno
            fonte = estilo.font
            fonte.name =  "Calibri (Corpo)"
            fonte.size = Pt(24)

    certificado = "C:\\Users\\nicho\\PycharmProjects\\Enviando_Email\\Certificados\\" + nome_aluno + ".docx"

    word.save(certificado)

print("Certificados gerados com sucesso!")
